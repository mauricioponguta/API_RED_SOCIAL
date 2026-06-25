from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from PIL import Image


CAPTURES_DIR = Path(
    r"C:\Users\MAURICIO\OneDrive\Documentos\ADSO\nest_red_social\Capturas de pantalla"
)
OUTPUT = Path("Documentacion_capturas_red_social.docx")


def set_run_font(run, size=None, bold=False, color=None):
    run.font.name = "Calibri"
    if size:
        run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def configure_styles(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color in [
        ("Heading 1", 16, (46, 116, 181)),
        ("Heading 2", 13, (46, 116, 181)),
        ("Heading 3", 12, (31, 77, 120)),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(*color)
        style.paragraph_format.space_before = Pt(14)
        style.paragraph_format.space_after = Pt(7)


def add_title(doc, images):
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("Documentacion de capturas - API Red Social")
    set_run_font(title_run, size=18, bold=True, color=(31, 77, 120))

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run(
        f"Registro visual de {len(images)} capturas de pantalla del proyecto."
    )
    set_run_font(subtitle_run, size=11, color=(85, 85, 85))

    doc.add_paragraph()
    summary = doc.add_paragraph()
    summary.add_run("Contenido: ").bold = True
    summary.add_run(
        "cada captura se presenta con su nombre de archivo para facilitar la revision y trazabilidad."
    )

    doc.add_page_break()


def image_width_for_page(image_path):
    with Image.open(image_path) as img:
        width, height = img.size

    max_width = 6.5
    max_height = 7.4
    ratio = width / height

    width_by_height = max_height * ratio
    return Inches(min(max_width, width_by_height))


def add_capture(doc, image_path, index, total):
    heading = doc.add_heading(f"Captura {index} de {total}", level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT

    meta = doc.add_paragraph()
    run = meta.add_run(image_path.name)
    set_run_font(run, size=9, color=(85, 85, 85))

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=image_width_for_page(image_path))

    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_run = caption.add_run(f"Figura {index}. {image_path.stem}")
    set_run_font(caption_run, size=9, color=(85, 85, 85))

    if index != total:
        doc.add_page_break()


def build():
    images = sorted(
        [
            path
            for path in CAPTURES_DIR.iterdir()
            if path.suffix.lower() in {".png", ".jpg", ".jpeg"}
        ],
        key=lambda p: p.name,
    )

    if not images:
        raise SystemExit(f"No se encontraron imagenes en {CAPTURES_DIR}")

    doc = Document()
    configure_styles(doc)
    add_title(doc, images)

    for index, image_path in enumerate(images, start=1):
        add_capture(doc, image_path, index, len(images))

    doc.core_properties.title = "Documentacion de capturas - API Red Social"
    doc.core_properties.subject = "Capturas de pantalla del proyecto"
    doc.core_properties.author = "Codex"
    doc.save(OUTPUT)
    print(OUTPUT.resolve())


if __name__ == "__main__":
    build()
